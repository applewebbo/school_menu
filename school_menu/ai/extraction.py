"""
Turning an uploaded file into rows ready for the preview (#234).

Order matters here: everything that can fail for free (extension, decoding, workbook
parsing) happens before the call, so a file we cannot use never costs a quota slot or a
token. After the call the answer is re-validated locally and cleaned deterministically,
so a bad answer becomes an error the user can act on, never a bad row in the database.
"""

import base64
import csv
import io
from dataclasses import dataclass

from django.conf import settings
from openpyxl import load_workbook
from pydantic import ValidationError

from school_menu.ai import prompts, schemas
from school_menu.ai.client import get_client
from school_menu.ai.errors import (
    ExtractionTimeout,
    InvalidResponse,
    UnsupportedFile,
    UpstreamError,
)
from school_menu.ai.normalise import normalise_rows

PDF_EXTENSIONS = {".pdf"}
TEXT_EXTENSIONS = {".csv", ".txt"}
SHEET_EXTENSIONS = {".xlsx"}
SUPPORTED_EXTENSIONS = PDF_EXTENSIONS | TEXT_EXTENSIONS | SHEET_EXTENSIONS

# Tried in order: most menus are UTF-8, files exported from Excel carry a BOM, and older
# ones come out of Windows as latin-1. The last one never fails, so decoding always works.
ENCODINGS = ("utf-8-sig", "utf-8", "latin-1")

# Only these are worth spending a second call on: a schema violation would come back the
# same way, so retrying it is paying twice for the same answer.
TRANSIENT_ERRORS = (UpstreamError, ExtractionTimeout)


@dataclass
class ExtractionResult:
    rows: list
    warnings: list
    usage: dict


def _extension(filename):
    _, _, extension = filename.lower().rpartition(".")
    return f".{extension}" if extension else ""


def _decode(content):
    for encoding in ENCODINGS:
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return ""  # pragma: no cover - latin-1 decodes any byte sequence


def _workbook_to_text(content):
    """
    Flatten every sheet to CSV.

    tablib reads only the first sheet, and school menus routinely use one sheet per week,
    so the missing sheets would silently disappear from the imported menu.
    """
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise UnsupportedFile(f"xlsx illeggibile: {exc}") from exc

    output = io.StringIO()
    writer = csv.writer(output)
    for sheet in workbook.worksheets:
        output.write(f"# {sheet.title}\n")
        for row in sheet.iter_rows(values_only=True):
            writer.writerow(["" if cell is None else cell for cell in row])
    workbook.close()
    return output.getvalue()


def _build_contents(content, filename):
    """Prepare what gets sent upstream, or refuse the file before spending anything."""
    extension = _extension(filename)
    if extension in PDF_EXTENSIONS:
        # Inline rather than the Files API: nothing is retained by Google beyond the
        # request itself, which is what the privacy policy promises. The payload has to
        # be base64 already — `data` is a Base64EncodedString, and raw bytes pass
        # through its validator untouched only to fail later as invalid UTF-8.
        return [
            {
                "type": "document",
                "mime_type": "application/pdf",
                "data": base64.b64encode(content).decode("ascii"),
            },
        ]
    if extension in TEXT_EXTENSIONS:
        text = _decode(content)
    elif extension in SHEET_EXTENSIONS:
        text = _workbook_to_text(content)
    else:
        raise UnsupportedFile(f"estensione non supportata: {extension or filename}")
    return [{"type": "text", "text": text[: settings.AI_MENU_IMPORT_MAX_TEXT_CHARS]}]


def _ask(contents, system_instruction, schema):
    """Call the model, retrying only what is worth retrying."""
    client = get_client()
    attempts = settings.AI_MENU_IMPORT_MAX_RETRIES + 1
    attempt = 0
    while True:
        attempt += 1
        try:
            return client.extract(
                system_instruction=system_instruction,
                contents=contents,
                schema=schema,
            )
        except TRANSIENT_ERRORS:
            if attempt >= attempts:
                raise


def extract_menu(kind, content, filename):
    """
    Extract the menu contained in an uploaded file.

    Args:
        kind: MenuImportDraft.Kinds value
        content: raw bytes of the uploaded file
        filename: original name, used only to pick the right preparation

    Returns:
        ExtractionResult with normalised rows, warnings for the user, and token usage.

    Raises:
        AIImportError: any failure, already carrying an Italian message for the user.
    """
    contents = _build_contents(content, filename)
    response = _ask(
        contents, prompts.build_system_instruction(kind), schemas.json_schema(kind)
    )

    try:
        raw_rows = schemas.parse(kind, response.text)
    except ValidationError as exc:
        raise InvalidResponse(str(exc)) from exc

    rows, warnings = normalise_rows(kind, raw_rows)
    return ExtractionResult(rows=rows, warnings=warnings, usage=response.usage)
